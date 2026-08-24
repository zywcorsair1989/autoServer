"""Document processing utilities.

This module provides functionality for processing different document
types and extracting text content for the RAG system.
"""

import os
import logging
from typing import List, Optional, Tuple
from dataclasses import dataclass

from app.utils.text_splitter import TextSplitter, TextChunk

logger = logging.getLogger(__name__)


@dataclass
class ProcessedDocument:
    """Represents a processed document with its chunks."""

    filename: str
    content: str
    chunks: List[TextChunk]
    metadata: dict


class DocumentProcessor:
    """Document processor for extracting and chunking text from files.

    Supports:
    - Plain text files (.txt)
    - PDF files (.pdf)
    - Word documents (.docx)
    """

    SUPPORTED_EXTENSIONS = {'.txt', '.pdf', '.docx'}

    def __init__(
        self,
        chunk_size: int = 500,
        chunk_overlap: int = 50
    ):
        """Initialize the document processor.

        Args:
            chunk_size: Target size for each chunk in characters.
            chunk_overlap: Number of characters to overlap between chunks.
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = TextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap
        )

    def process_file(
        self,
        file_path: str,
        metadata: Optional[dict] = None
    ) -> ProcessedDocument:
        """Process a file and return its processed content.

        Args:
            file_path: Path to the file.
            metadata: Optional metadata to include.

        Returns:
            ProcessedDocument with content and chunks.

        Raises:
            ValueError: If file type is not supported.
            FileNotFoundError: If file does not exist.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        # Get file extension
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {ext}. "
                f"Supported types: {', '.join(self.SUPPORTED_EXTENSIONS)}"
            )

        # Extract text based on file type
        if ext == '.txt':
            content = self._extract_text_from_txt(file_path)
        elif ext == '.pdf':
            content = self._extract_text_from_pdf(file_path)
        elif ext == '.docx':
            content = self._extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        # Chunk the content
        chunks = self.text_splitter.split_text(content)

        # Create metadata
        filename = os.path.basename(file_path)
        file_metadata = {
            'filename': filename,
            'file_path': file_path,
            'file_size': os.path.getsize(file_path),
            'file_type': ext,
        }
        if metadata:
            file_metadata.update(metadata)

        logger.info(
            f"Processed {filename}: {len(content)} chars, "
            f"{len(chunks)} chunks"
        )

        return ProcessedDocument(
            filename=filename,
            content=content,
            chunks=chunks,
            metadata=file_metadata
        )

    def _extract_text_from_txt(self, file_path: str) -> str:
        """Extract text from a plain text file.

        Args:
            file_path: Path to the text file.

        Returns:
            Extracted text content.
        """
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

    def _extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from a PDF file.

        Args:
            file_path: Path to the PDF file.

        Returns:
            Extracted text content.
        """
        try:
            import pypdf
        except ImportError:
            logger.warning("pypdf not installed, trying PyPDF2")
            try:
                import PyPDF2 as pypdf
            except ImportError:
                raise ImportError(
                    "PDF extraction requires pypdf or PyPDF2. "
                    "Install with: pip install pypdf"
                )

        text_parts = []

        with open(file_path, 'rb') as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

        return "\n\n".join(text_parts)

    def _extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from a Word document.

        Args:
            file_path: Path to the .docx file.

        Returns:
            Extracted text content.
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "Word document extraction requires python-docx. "
                "Install with: pip install python-docx"
            )

        doc = Document(file_path)
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        return "\n\n".join(text_parts)

    def process_text(
        self,
        text: str,
        filename: str = "text",
        metadata: Optional[dict] = None
    ) -> ProcessedDocument:
        """Process plain text content.

        Args:
            text: Text content to process.
            filename: Virtual filename for the content.
            metadata: Optional metadata to include.

        Returns:
            ProcessedDocument with content and chunks.
        """
        chunks = self.text_splitter.split_text(text)

        file_metadata = {
            'filename': filename,
            'file_type': '.txt',
        }
        if metadata:
            file_metadata.update(metadata)

        return ProcessedDocument(
            filename=filename,
            content=text,
            chunks=chunks,
            metadata=file_metadata
        )

    @classmethod
    def get_supported_extensions(cls) -> List[str]:
        """Get list of supported file extensions.

        Returns:
            List of supported file extensions.
        """
        return list(cls.SUPPORTED_EXTENSIONS)

    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        """Check if a file type is supported.

        Args:
            file_path: Path to the file.

        Returns:
            True if file type is supported, False otherwise.
        """
        _, ext = os.path.splitext(file_path)
        return ext.lower() in cls.SUPPORTED_EXTENSIONS